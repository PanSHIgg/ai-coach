#!/usr/bin/env python3
"""gen_calorie_cases.py — DATA4 热量对拍用例生成（50 行）。

用 backend 的 app.services.calorie.compute_account 计算期望值，导出：
- data/testcases/calorie_test_cases.xlsx（openpyxl 可用时）
- 如 openpyxl 不可用，则同时导出 calorie_test_cases.csv

覆盖：5 种活动水平 × 男/女 × 5 种（摄入/运动/步数）场景，含 steps=0 与 steps>0，
含 hourly（逐时步频）与不含 hourly 两种校准路径。

口径（与引擎完全同源，误差 0）：
- BMR = Mifflin-St Jeor（男 +5 / 女 −161），age 由 birth_year 按 REF_YEAR 推算；
- PAL = PAL_BY_LEVEL[activity_level]；
- TEF = intake × 10%；
- TEE = BMR×PAL + eat + TEF；
- deficit = intake − TEE；
- today_left = BMR×PAL − intake + eat − TEF。

【假设】age 依赖年份：本脚本以 REF_YEAR=2026 为基准由 birth_year 估算 age（与引擎
age_from_birth_year 的「当年年份 − birth_year」口径一致）。若跨年复跑，需同步更新 REF_YEAR
并重新生成，以保证期望值与引擎口径仍误差 0。
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

# 引入后端热量引擎（不修改 backend 任何文件，仅只读复用）
BACKEND_DIR = Path(__file__).resolve().parents[2] / "backend"
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from app.services.calorie import CalorieInputs, compute_account  # noqa: E402

OUT_DIR = Path(__file__).resolve().parent
OUT_XLSX = OUT_DIR / "calorie_test_cases.xlsx"
OUT_CSV = OUT_DIR / "calorie_test_cases.csv"

REF_YEAR = 2026  # 见模块 docstring【假设】

HEADERS = [
    "case_id", "gender", "height_cm", "weight_kg", "birth_year", "activity_level",
    "intake", "eat", "steps", "hourly_json",
    "expected_bmr", "expected_pal", "expected_tef", "expected_tee",
    "expected_deficit", "expected_today_left",
]

# 5 种场景：(intake, eat, steps, hourly)
SCENARIOS = [
    (1500, 0, 0, None),
    (1800, 300, 6000, None),
    (1200, 0, 10000, None),
    (2000, 500, 12000, {"8": 300, "9": 600, "12": 1500, "18": 2200, "20": 1200}),
    (1600, 200, 15000,
     {"7": 200, "8": 500, "9": 800, "12": 1800, "18": 3000, "19": 2500, "20": 1500}),
]


def build_rows() -> list[list]:
    rows: list[list] = []
    n = 0
    for level in range(1, 6):
        for gender in (1, 2):  # 1 男 / 2 女
            if gender == 1:
                height = 175.0 + (level - 3) * 2.0
                weight = 75.0 + (level - 3) * 3.0
                birth_year = 1990 + level
            else:
                height = 163.0 + (level - 3) * 2.0
                weight = 58.0 + (level - 3) * 2.0
                birth_year = 1992 + level
            age = REF_YEAR - birth_year
            for intake, eat, steps, hourly in SCENARIOS:
                n += 1
                res = compute_account(CalorieInputs(
                    gender=gender, weight_kg=weight, height_cm=height, age=age,
                    activity_level=level, intake=intake, eat=eat, steps=steps, hourly=hourly,
                ))
                rows.append([
                    f"CC-{n:03d}", gender, height, weight, birth_year, level,
                    float(intake), float(eat), int(steps),
                    json.dumps(hourly, ensure_ascii=False, sort_keys=True) if hourly else "",
                    res.bmr, res.pal, res.tef, res.tee, res.deficit, res.today_left,
                ])
    return rows


def write_xlsx(rows: list[list]) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "calorie_cases"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    # 列宽（可读性）
    widths = [9, 7, 11, 11, 11, 14, 8, 8, 8, 34, 12, 11, 11, 11, 15, 17]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    wb.save(OUT_XLSX)
    return True


def write_csv(rows: list[list]) -> None:
    import csv
    with open(OUT_CSV, "w", newline="", encoding="utf-8-sig") as fp:
        w = csv.writer(fp)
        w.writerow(HEADERS)
        w.writerows(rows)


def main() -> int:
    rows = build_rows()
    assert len(rows) == 50, f"用例数应为 50，实际 {len(rows)}"
    ok = write_xlsx(rows)
    if ok:
        print(f"calorie_test_cases.xlsx: {len(rows)} 行 -> {OUT_XLSX}")
    else:
        write_csv(rows)
        print("openpyxl 不可用，已改导出 CSV：")
        print(f"calorie_test_cases.csv: {len(rows)} 行 -> {OUT_CSV}")
    # 打印 2 行样例便于核对
    for r in (rows[0], rows[24]):
        print(r)
    return 0


if __name__ == "__main__":
    sys.exit(main())